"""
storage.py
Слой хранения данных приложения "Бюджет семьи".
Все данные хранятся локально в SQLite (файл budget.db в папке приложения),
поэтому приложение полностью работает офлайн.
"""

import sqlite3
import datetime
import os


class Storage:
    def __init__(self, db_path: str):
        self.db_path = db_path
        os.makedirs(os.path.dirname(db_path), exist_ok=True)
        self._conn = sqlite3.connect(self.db_path)
        self._conn.row_factory = sqlite3.Row
        self._init_db()

    def _init_db(self):
        cur = self._conn.cursor()
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS settings (
                key   TEXT PRIMARY KEY,
                value TEXT
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS expenses (
                id       INTEGER PRIMARY KEY AUTOINCREMENT,
                date     TEXT NOT NULL,
                amount   REAL NOT NULL,
                category TEXT NOT NULL
            )
            """
        )
        self._conn.commit()

        # значения по умолчанию, если ещё не заданы
        if self.get_setting("income_husband") is None:
            self.set_setting("income_husband", "0")
        if self.get_setting("income_wife") is None:
            self.set_setting("income_wife", "0")

    # ---------- настройки / доходы ----------

    def get_setting(self, key, default=None):
        cur = self._conn.execute("SELECT value FROM settings WHERE key = ?", (key,))
        row = cur.fetchone()
        return row["value"] if row else default

    def set_setting(self, key, value):
        self._conn.execute(
            "INSERT INTO settings (key, value) VALUES (?, ?) "
            "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
            (key, str(value)),
        )
        self._conn.commit()

    def get_incomes(self):
        husband = float(self.get_setting("income_husband", "0") or 0)
        wife = float(self.get_setting("income_wife", "0") or 0)
        return husband, wife

    def set_incomes(self, husband: float, wife: float):
        self.set_setting("income_husband", husband)
        self.set_setting("income_wife", wife)

    # ---------- расходы ----------

    def add_expense(self, date: str, amount: float, category: str):
        self._conn.execute(
            "INSERT INTO expenses (date, amount, category) VALUES (?, ?, ?)",
            (date, amount, category),
        )
        self._conn.commit()

    def delete_expense(self, expense_id: int):
        self._conn.execute("DELETE FROM expenses WHERE id = ?", (expense_id,))
        self._conn.commit()

    def list_expenses(self, limit=None):
        query = "SELECT id, date, amount, category FROM expenses ORDER BY date DESC, id DESC"
        if limit:
            query += f" LIMIT {int(limit)}"
        cur = self._conn.execute(query)
        return cur.fetchall()

    def total_expenses(self):
        cur = self._conn.execute("SELECT COALESCE(SUM(amount), 0) AS total FROM expenses")
        return cur.fetchone()["total"]

    def top_categories(self, limit=6):
        """Самые часто используемые категории — для быстрых кнопок ввода."""
        cur = self._conn.execute(
            """
            SELECT category, COUNT(*) AS cnt
            FROM expenses
            GROUP BY category
            ORDER BY cnt DESC
            LIMIT ?
            """,
            (limit,),
        )
        return [row["category"] for row in cur.fetchall()]

    # ---------- сводка ----------

    def summary(self):
        husband, wife = self.get_incomes()
        total_income = husband + wife
        total_expenses = self.total_expenses()
        remaining = total_income - total_expenses
        return {
            "husband": husband,
            "wife": wife,
            "total_income": total_income,
            "total_expenses": total_expenses,
            "remaining": remaining,
        }

    # ---------- экспорт в xlsx (формат как в исходной таблице) ----------

    def export_to_xlsx(self, path: str):
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Лист1"

        ws["A1"] = "Доход мужа"
        ws["B1"] = self.get_incomes()[0]
        ws["A2"] = "Доход жены"
        ws["B2"] = self.get_incomes()[1]
        ws["A3"] = "Общий доход"
        ws["B3"] = "=B1+B2"
        ws["A4"] = "Расходы семьи"
        ws["A5"] = "Остаток денег"
        ws["B5"] = "=B3-B4"

        ws["A7"] = "Дата"
        ws["B7"] = "Сумма расхода"
        ws["C7"] = "Наименование расхода"

        rows = list(self.list_expenses())
        rows.reverse()  # в исходной таблице расходы шли от старых к новым
        row_idx = 8
        for r in rows:
            date_val = datetime.date.fromisoformat(r["date"])
            ws.cell(row=row_idx, column=1, value=date_val)
            ws.cell(row=row_idx, column=1).number_format = "DD.MM.YYYY"
            ws.cell(row=row_idx, column=2, value=r["amount"])
            ws.cell(row=row_idx, column=3, value=r["category"])
            row_idx += 1

        last_row = row_idx - 1
        if last_row >= 8:
            ws["B4"] = f"=SUM(B8:B{last_row})"
        else:
            ws["B4"] = 0

        for col, width in (("A", 14), ("B", 16), ("C", 24)):
            ws.column_dimensions[col].width = width

        wb.save(path)
        return path

    def close(self):
        self._conn.close()
